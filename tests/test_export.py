from datetime import datetime

import openpyxl

from conftest import make_report
from handover.export import SHEETS, export_workbook
from handover.insights import InsightPoint, Insights, InsightSnapshot
from handover.schema import Contact, Item, TimelineEntry, Unconfirmed
from handover.store import Record

NOW = datetime(2026, 10, 15, 10, 0)


def rec(tmp_path, name, priority, category, warnings=(), **kw):
    report = make_report(priority=priority, category=category, summary=f"{name} 요약",
                         priority_reason=f"{name} 이유", **kw)
    return Record(source_name=name, created_at=NOW, memo_hash=name, memo=f"{name} 원문",
                  report=report, warnings=list(warnings), txt_path=tmp_path / f"{name}.txt")


def sample(tmp_path):
    records = [
        rec(tmp_path, "a.txt", "일반", "정상·기록"),
        rec(tmp_path, "b.txt", "긴급", "설비 이상", warnings=["원문에 없는 숫자 '3': x"],
            events=[Item(text="2호기 누유", source_quote="q")],
            timeline=[TimelineEntry(time="07:50", text="액체 발견", source_quote="q")],
            contacts=[Contact(party="설비팀", request="점검", status="미회신", source_quote="q")],
            unconfirmed=[Unconfirmed(subject="원인", note="")]),
        rec(tmp_path, "c.txt", "주의", "수량 불일치", unconfirmed=[Unconfirmed(subject="수량", note="못 셈")]),
    ]
    insights = Insights(headline="설비 이상이 핵심", risks=[],
                        patterns=[InsightPoint(title="회신 대기", detail="두 건", related=["b.txt", "c.txt"])],
                        gaps=[InsightPoint(title="원인 미기록", detail="반복", related=["b.txt"])])
    snapshot = InsightSnapshot(insights, NOW, 3, [])
    return records, snapshot


def test_workbook_has_all_sheets(tmp_path):
    records, snap = sample(tmp_path)
    path = export_workbook(records, snap, tmp_path / "out.xlsx", NOW)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == SHEETS == ["대시보드", "처리현황", "상세보고서", "AI인사이트", "자동점검경고"]


def test_status_sheet_sorted_by_priority_with_filter(tmp_path):
    records, snap = sample(tmp_path)
    ws = openpyxl.load_workbook(export_workbook(records, snap, tmp_path / "o.xlsx", NOW))["처리현황"]
    header = [c.value for c in ws[1]]
    assert header[:5] == ["긴급도", "메모", "문제 유형", "요약", "판단 이유"]
    assert [ws.cell(row=r, column=1).value for r in (2, 3, 4)] == ["긴급", "주의", "일반"]
    assert ws.auto_filter.ref is not None and ws.freeze_panes == "A2"


def test_detail_sheet_one_row_per_item(tmp_path):
    records, snap = sample(tmp_path)
    ws = openpyxl.load_workbook(export_workbook(records, snap, tmp_path / "o.xlsx", NOW))["상세보고서"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    b_rows = [r for r in rows if r[0] == "b.txt"]
    assert ["b.txt", "긴급", "발생한 일", "2호기 누유"] == b_rows[0][:4] or any(
        r[2] == "발생한 일" and r[3] == "2호기 누유" for r in b_rows)
    assert any(r[2] == "시간순 경과" and r[3] == "07:50  액체 발견" for r in b_rows)
    assert any(r[2] == "미확인 사항" and r[3] == "원인: 미확인" for r in b_rows)
    assert any(r[2] == "요약" for r in b_rows)


def test_dashboard_sheet_counts_and_charts(tmp_path):
    records, snap = sample(tmp_path)
    ws = openpyxl.load_workbook(export_workbook(records, snap, tmp_path / "o.xlsx", NOW))["대시보드"]
    values = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}
    assert values["전체"] == 3 and values["긴급"] == 1 and values["경고"] == 1
    assert len(ws._charts) == 2  # 문제 유형 원형 차트 + 미확인 막대 차트


def test_insights_and_warnings_sheets(tmp_path):
    records, snap = sample(tmp_path)
    wb = openpyxl.load_workbook(export_workbook(records, snap, tmp_path / "o.xlsx", NOW))
    ins = [[c.value for c in row] for row in wb["AI인사이트"].iter_rows(min_row=1)]
    assert any(r[0] == "공통 패턴" and r[1] == "회신 대기" and r[3] == "b.txt, c.txt" for r in ins)
    warns = [[c.value for c in row] for row in wb["자동점검경고"].iter_rows(min_row=2)]
    assert warns == [["b.txt", "긴급", "원문에 없는 숫자 '3': x"]]


def test_export_with_every_pie_slice(tmp_path):
    cats = ["설비 이상", "수량 불일치", "품질·검사", "안전", "라벨·식별", "서류·전산", "미분류", "정상·기록"]
    records = [rec(tmp_path, f"{i}.txt", "주의", c, unconfirmed=[Unconfirmed(subject="x", note="")])
               for i, c in enumerate(cats * 2)]
    _, snap = sample(tmp_path)
    ws = openpyxl.load_workbook(export_workbook(records, snap, tmp_path / "o.xlsx", NOW))["대시보드"]
    assert len(ws._charts) == 2


def test_export_without_insights(tmp_path):
    records, _ = sample(tmp_path)
    wb = openpyxl.load_workbook(export_workbook(records, None, tmp_path / "o.xlsx", NOW))
    assert "분석한 인사이트가 없습니다" in str([c.value for row in wb["AI인사이트"].iter_rows() for c in row])
