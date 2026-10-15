import tkinter as tk

import pytest

from conftest import SAMPLE_REPORT, FakeClient, make_report
from handover.schema import Contact, Fact, Item, TimelineEntry, Unconfirmed
from handover.insights import InsightPoint, Insights
from handover.view import build_sections


def test_build_sections_order_and_placeholders():
    r = make_report(unconfirmed=[Unconfirmed(subject="라벨 여분 수량", note="수량은 못 셈."),
                                 Unconfirmed(subject="원인", note="")])
    sections = build_sections(r)
    assert [s.title for s in sections] == ["발생한 일", "조치한 내용", "현재 상태", "다음 근무자 확인 사항",
                                           "시간순 경과", "주요 수치·대상", "연락·요청 현황", "미확인 사항"]
    assert sections[0].lines == ["메모에 기록 없음"]
    assert sections[7].lines == ["라벨 여분 수량: 미확인 (수량은 못 셈.)", "원인: 미확인"]
    assert sections[7].kind == "unconfirmed"


def test_unconfirmed_ranking_orders_and_limits():
    from handover.view import unconfirmed_ranking
    u = lambda n: [Unconfirmed(subject=f"s{i}", note="") for i in range(n)]
    records = [type("R", (), {"source_name": name, "report": make_report(priority=p, unconfirmed=u(n))})()
               for name, p, n in [("a", "일반", 3), ("b", "긴급", 3), ("c", "주의", 5), ("d", "긴급", 0), ("e", "주의", 1)]]
    assert unconfirmed_ranking(records, limit=3) == [("c", "주의", 5), ("b", "긴급", 3), ("a", "일반", 3)]


def test_dashboard_chart_draws_one_bar_per_ranked_memo(root, tmp_path):
    write_memos(tmp_path, ["a", "b"])
    report = make_report(priority="긴급", unconfirmed=[Unconfirmed(subject="원인", note="")])
    app = make_app(root, tmp_path, FakeClient("end_turn", report))
    app.run_batch(); app.poll_results()
    assert app.chart_bars() == 2


def three_priority_app(root, tmp_path):
    write_memos(tmp_path, ["a", "b", "c"])
    reports = [make_report(priority="긴급", category="설비 이상"),
               make_report(priority="주의", category="수량 불일치"),
               make_report(priority="주의", category="설비 이상",
                           events=[Item(text="라벨 10장", source_quote="없는 문장")])]
    app = make_app(root, tmp_path, FakeClient(results=[("end_turn", r) for r in reports]))
    app.run_batch(); app.poll_results()
    return app


def test_tile_click_filters_table(root, tmp_path):
    app = three_priority_app(root, tmp_path)
    app.click_tile("주의")
    assert [r[0] for r in app.dashboard_rows()] == ["주의", "주의"]
    assert app.tile_counts()["전체"] == 3  # 타일 숫자는 필터와 상관없이 전체 기준
    app.click_tile("경고")
    assert len(app.dashboard_rows()) == 1 and app.dashboard_rows()[0][5] > 0
    app.click_tile("경고")  # 같은 타일 다시 누르면 해제
    assert len(app.dashboard_rows()) == 3
    app.click_tile("긴급"); app.click_tile("전체")
    assert len(app.dashboard_rows()) == 3


def test_pie_slices_and_slice_filter(root, tmp_path):
    app = three_priority_app(root, tmp_path)
    assert app.pie_slices() == [("설비 이상", 2), ("수량 불일치", 1)]
    app.filter_category("수량 불일치")
    assert [r[1] for r in app.dashboard_rows()] == ["c.txt"] or len(app.dashboard_rows()) == 1
    assert "수량 불일치" in app.filter_text()


def test_unclassified_records_are_classified_before_insights(root, tmp_path):
    import json
    write_memos(tmp_path, ["a"])
    app = make_app(root, tmp_path, FakeClient("end_turn", SAMPLE_REPORT))
    app.run_batch(); app.poll_results()
    path = next((tmp_path / "reports").glob("인수인계_*.json"))
    data = json.loads(path.read_text(encoding="utf-8"))
    del data["report"]["category"]
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    from handover.classify import Classification, MemoCategory
    fake = FakeClient("end_turn", SAMPLE_REPORT,
                      classification=Classification(items=[MemoCategory(memo="a.txt", category="안전")]))
    window = tk.Toplevel(root)
    try:
        again = make_app(window, tmp_path, fake)
        assert again.pie_slices() == [("미분류", 1)]
        again.run_insights(); again.poll_results()
        assert len(fake.classify_calls) == 1
        assert again.pie_slices() == [("안전", 1)]
    finally:
        window.destroy()


def test_export_excel_writes_all_sheets(root, tmp_path):
    import openpyxl
    app = three_priority_app(root, tmp_path)
    app.click_tile("긴급")  # 필터와 상관없이 전체를 내보낸다
    path = app.export_excel(tmp_path / "out.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:2] == ["대시보드", "처리현황"]
    assert wb["처리현황"].max_row == 4
    assert "엑셀로 내보냈습니다" in app.batch_status_text()


def test_export_excel_without_records(root, tmp_path):
    app = make_app(root, tmp_path, FakeClient("end_turn", SAMPLE_REPORT))
    assert app.export_excel(tmp_path / "out.xlsx") is None
    assert "내보낼 보고서가 없습니다" in app.batch_status_text()
    assert not (tmp_path / "out.xlsx").exists()


def test_build_sections_detail_formats():
    r = make_report(
        timeline=[TimelineEntry(time="14:10", text="진동 발생", source_quote="q")],
        key_facts=[Fact(label="대기 제품", value="12개", source_quote="q")],
        contacts=[Contact(party="설비팀", request="점검 요청", status="결과 미회신", source_quote="q")],
    )
    by_kind = {s.kind: s.lines for s in build_sections(r)}
    assert by_kind["timeline"] == ["14:10  진동 발생"]
    assert by_kind["facts"] == ["대기 제품: 12개"]
    assert by_kind["contacts"] == ["설비팀 — 점검 요청 — 결과 미회신"]


@pytest.fixture(scope="session")
def tk_session():
    # Tk 인터프리터를 테스트마다 새로 만들면 Windows에서 가끔 초기화가 실패한다 → 세션당 하나만 만든다.
    try:
        r = tk.Tk()
    except tk.TclError:
        pytest.skip("화면(Tk)을 사용할 수 없는 환경")
    r.withdraw()
    yield r
    r.destroy()


@pytest.fixture
def root(tk_session):
    window = tk.Toplevel(tk_session)
    window.withdraw()
    yield window
    window.destroy()


def make_app(root, tmp_path, client):
    from app import HandoverApp
    return HandoverApp(root, client_factory=lambda: client, out_dir=tmp_path / "reports",
                       samples_dir=tmp_path, runner=lambda fn: fn())


def test_app_generates_saves_and_lists_history(root, tmp_path):
    fake = FakeClient("end_turn", SAMPLE_REPORT)
    app = make_app(root, tmp_path, fake)
    app.load_memo_text("08:40 라벨 비뚤어짐.", "작업메모_01")
    app.generate()
    app.poll_results()
    saved = list((tmp_path / "reports").glob("인수인계_작업메모_01_*.txt"))
    assert len(saved) == 1
    assert app.history_paths == saved
    assert "이상 없음" in app.status_text()


def test_app_blank_memo_does_not_call_api(root, tmp_path):
    fake = FakeClient("end_turn", SAMPLE_REPORT)
    app = make_app(root, tmp_path, fake)
    app.generate()
    assert fake.calls == []
    assert "메모를 입력" in app.status_text()


def test_app_shows_error_on_failure(root, tmp_path):
    app = make_app(root, tmp_path, FakeClient("refusal", None))
    app.load_memo_text("메모", "m")
    app.generate()
    app.poll_results()
    assert "거절" in app.status_text()
    assert not (tmp_path / "reports").exists() or not list((tmp_path / "reports").iterdir())


def test_app_starts_on_dashboard(root, tmp_path):
    app = make_app(root, tmp_path, FakeClient("end_turn", SAMPLE_REPORT))
    assert app.current_tab == "dashboard"


def test_single_report_syncs_to_dashboard(root, tmp_path):
    write_memos(tmp_path, ["a", "b"])
    fake = FakeClient("end_turn", SAMPLE_REPORT)
    app = make_app(root, tmp_path, fake)
    app.run_batch(); app.poll_results()
    assert app.tile_counts()["전체"] == 2
    app.show_tab("editor")
    app.load_memo_text("08:40 라벨 비뚤어짐.", "직접 입력")
    app.generate(); app.poll_results()
    app.show_tab("dashboard")
    assert app.tile_counts()["전체"] == 3
    assert "직접 입력" in [r[1] for r in app.dashboard_rows()]


def test_typed_memos_are_not_merged_on_dashboard(root, tmp_path):
    app = make_app(root, tmp_path, FakeClient("end_turn", SAMPLE_REPORT))
    for memo in ("08:40 라벨 비뚤어짐.", "09:10 센서 청소 후 재가동."):
        app.clear_memo()
        app.memo.insert("1.0", memo)
        app.generate(); app.poll_results()
    assert app.tile_counts()["전체"] == 2


def write_memos(folder, names):
    for name in names:
        (folder / f"{name}.txt").write_text(f"{name} 메모", encoding="utf-8")


def test_dashboard_batch_sorted_and_counted(root, tmp_path):
    write_memos(tmp_path, ["a", "b", "c"])
    fake = FakeClient(results=[("end_turn", make_report(priority=p, summary=p)) for p in ("일반", "긴급", "주의")])
    app = make_app(root, tmp_path, fake)
    app.show_tab("dashboard")
    app.run_batch()
    app.poll_results()
    rows = app.dashboard_rows()
    assert [r[0] for r in rows] == ["긴급", "주의", "일반"]
    assert app.tile_counts() == {"전체": 3, "긴급": 1, "주의": 1, "일반": 1, "경고": 0, "실패": 0}


def test_dashboard_second_batch_skips_processed(root, tmp_path):
    write_memos(tmp_path, ["a", "b"])
    fake = FakeClient("end_turn", SAMPLE_REPORT)
    app = make_app(root, tmp_path, fake)
    app.run_batch()
    app.poll_results()
    app.run_batch()
    app.poll_results()
    assert len(fake.calls) == 2
    assert len(app.dashboard_rows()) == 2


def test_dashboard_empty_folder(root, tmp_path):
    fake = FakeClient("end_turn", SAMPLE_REPORT)
    app = make_app(root, tmp_path, fake)
    app.run_batch()
    assert ".txt 메모가 없습니다" in app.batch_status_text()
    assert fake.calls == []


def test_dashboard_ignores_second_run_while_busy(root, tmp_path):
    write_memos(tmp_path, ["a"])
    jobs = []
    from app import HandoverApp
    app = HandoverApp(root, client_factory=lambda: FakeClient("end_turn", SAMPLE_REPORT),
                      out_dir=tmp_path / "reports", samples_dir=tmp_path, runner=jobs.append)
    app.run_batch()
    app.run_batch()
    assert len(jobs) == 1


def test_dashboard_select_shows_detail(root, tmp_path):
    write_memos(tmp_path, ["a"])
    report = make_report(summary="2호기 누유 의심", priority="긴급", priority_reason="설비 아래 액체 발견")
    app = make_app(root, tmp_path, FakeClient("end_turn", report))
    app.run_batch()
    app.poll_results()
    app.select_dashboard_row(0)
    assert "설비 아래 액체 발견" in app.detail_status_text()


def test_dashboard_loads_existing_records_on_start(root, tmp_path):
    write_memos(tmp_path, ["a"])
    app = make_app(root, tmp_path, FakeClient("end_turn", SAMPLE_REPORT))
    app.run_batch()
    app.poll_results()
    window = tk.Toplevel(root)
    try:
        again = make_app(window, tmp_path, FakeClient("end_turn", SAMPLE_REPORT))
        assert len(again.dashboard_rows()) == 1
    finally:
        window.destroy()


INSIGHTS = Insights(headline="설비 이상 건이 회신 대기 중", patterns=[], risks=[],
                    gaps=[InsightPoint(title="원인 미기록", detail="원인이 반복 미확인", related=["a.txt"])])


def test_batch_then_insights_shown(root, tmp_path):
    write_memos(tmp_path, ["a", "b"])
    fake = FakeClient(results=[("end_turn", SAMPLE_REPORT), ("end_turn", SAMPLE_REPORT)], insights=INSIGHTS)
    app = make_app(root, tmp_path, fake)
    app.run_batch()
    app.poll_results()   # 일괄 처리 완료 → 인사이트 분석 시작
    app.poll_results()   # 인사이트 결과 반영
    text = app.insights_text()
    assert "설비 이상 건이 회신 대기 중" in text and "원인 미기록" in text
    assert (tmp_path / "reports" / "insights.json").is_file()


def test_insights_without_records(root, tmp_path):
    fake = FakeClient("end_turn", INSIGHTS)
    app = make_app(root, tmp_path, fake)
    app.run_insights()
    assert "처리된 보고서가 없습니다" in app.insights_text()
    assert fake.insight_calls == []


def test_saved_insights_loaded_and_marked_stale(root, tmp_path):
    write_memos(tmp_path, ["a"])
    fake = FakeClient(results=[("end_turn", SAMPLE_REPORT)], insights=INSIGHTS)
    app = make_app(root, tmp_path, fake)
    app.run_batch(); app.poll_results(); app.poll_results()
    assert "다시 분석하세요" not in app.insights_text()

    # 분석 이후 보고서가 하나 더 생김 → 저장된 인사이트는 보이되 '오래됨' 표시
    from datetime import datetime
    from handover.store import Record, save_record
    save_record(Record("b.txt", datetime(2026, 10, 15, 12, 0), "hb", "m", make_report(),
                       [], tmp_path / "reports" / "인수인계_b_20261015_120000.txt"))
    window = tk.Toplevel(root)
    try:
        again = make_app(window, tmp_path, FakeClient("end_turn", INSIGHTS))
        text = again.insights_text()
        assert "설비 이상 건이 회신 대기 중" in text and "다시 분석하세요" in text
    finally:
        window.destroy()
