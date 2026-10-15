"""대시보드·처리현황·상세보고서·AI 인사이트·자동점검경고를 엑셀 파일 하나로 내보낸다."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from handover.insights import InsightSnapshot
from handover.schema import PRIORITY_ORDER
from handover.store import Record, sort_for_dashboard
from handover.view import NOT_A_PROBLEM, other_breakdown, pie_groups, unconfirmed_ranking

SHEETS = ["대시보드", "처리현황", "상세보고서", "AI인사이트", "자동점검경고"]

# 앱과 같은 색 (긴급도 옅은 배경, 문제 유형 고정 색)
PRIORITY_FILL = {"긴급": "FEF2F2", "주의": "FFFBEB", "일반": "F0FDF4"}
PRIORITY_INK = {"긴급": "B91C1C", "주의": "B45309", "일반": "047857"}
STATUS_BAR = {"긴급": "D03B3B", "주의": "EC835A", "일반": "0CA30C"}
PIE_COLORS = {"설비 이상": "2A78D6", "수량 불일치": "EB6834", "품질·검사": "1BAF7A", "안전": "EDA100",
              "라벨·식별": "E87BA4", "기타": "8A8A85", "미분류": "C9C8C2"}

HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16)
SECTION_FONT = Font(bold=True, size=12)
THIN = Border(bottom=Side(style="thin", color="E2E8F0"))
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws, row: int, values: list[str], col: int = 1):
    for i, value in enumerate(values):
        cell = ws.cell(row=row, column=col + i, value=value)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def _widths(ws, widths: dict[str, int]):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _table(ws, headers: list[str], rows: list[list], widths: dict[str, int], priority_col: int | None = None):
    _header(ws, 1, headers)
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment, cell.border = WRAP, THIN
        if priority_col is not None:
            priority = values[priority_col - 1]
            if priority in PRIORITY_FILL:
                for c in range(1, len(headers) + 1):
                    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=PRIORITY_FILL[priority])
                ws.cell(row=r, column=priority_col).font = Font(bold=True, color=PRIORITY_INK[priority])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    _widths(ws, widths)


def _dashboard(ws, records: list[Record], snapshot: InsightSnapshot | None, now: datetime):
    ws["A1"] = "교대 인수인계 대시보드"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"내보낸 시각 {now:%Y-%m-%d %H:%M} · 보고서 {len(records)}건"
    ws["A2"].font = Font(color="64748B")

    ws["A4"] = "요약"
    ws["A4"].font = SECTION_FONT
    counts = {p: sum(1 for r in records if r.report.priority == p) for p in PRIORITY_ORDER}
    summary = [("전체", len(records)), *counts.items(), ("경고", sum(1 for r in records if r.warnings))]
    for i, (key, value) in enumerate(summary, start=5):
        ws.cell(row=i, column=1, value=key).font = Font(bold=True, color=PRIORITY_INK.get(key, "0F172A"))
        ws.cell(row=i, column=2, value=value)
    row = 5 + len(summary) + 1

    ws.cell(row=row, column=1, value="AI 인사이트").font = SECTION_FONT
    headline = snapshot.insights.headline if snapshot else "분석한 인사이트가 없습니다"
    ws.cell(row=row + 1, column=1, value=headline).alignment = WRAP
    # A:B만 병합한다 — D열부터는 유형 비율·미확인 표가 있어 겹치면 안 된다
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=2)
    ws.row_dimensions[row + 1].height = 120

    # 문제 유형 비율 표 + 원형 차트
    ws["D4"] = "문제 유형 비율 (정상·기록 제외)"
    ws["D4"].font = SECTION_FONT
    _header(ws, 5, ["유형", "건수", "비율"], col=4)
    slices = pie_groups(records)
    total = sum(n for _, n in slices) or 1
    for i, (name, count) in enumerate(slices, start=6):
        ws.cell(row=i, column=4, value=name)
        ws.cell(row=i, column=5, value=count)
        ws.cell(row=i, column=6, value=count / total).number_format = "0%"
    normal = sum(1 for r in records if r.report.category == NOT_A_PROBLEM)
    note_row = 6 + len(slices)
    ws.cell(row=note_row, column=4, value=f"정상·기록 {normal}건 제외").font = Font(color="94A3B8")
    other = other_breakdown(records)
    if other:
        ws.cell(row=note_row + 1, column=4, value="기타: " + ", ".join(f"{c} {n}건" for c, n in other)).font = \
            Font(color="94A3B8")
    if slices:
        pie = PieChart()
        pie.title = "문제 유형 비율"
        pie.add_data(Reference(ws, min_col=5, min_row=5, max_row=5 + len(slices)), titles_from_data=True)
        pie.set_categories(Reference(ws, min_col=4, min_row=6, max_row=5 + len(slices)))
        for i, (name, _) in enumerate(slices):
            point = DataPoint(idx=i)
            point.graphicalProperties.solidFill = PIE_COLORS.get(name, "8A8A85")
            point.graphicalProperties.line.solidFill = "FFFFFF"  # 조각 사이 흰 틈
            pie.series[0].dPt.append(point)
        pie.height, pie.width = 8, 12
        ws.add_chart(pie, "H4")

    # 미확인 항목이 많은 메모 표 + 막대 차트
    start = max(note_row + 3, 16)
    ws.cell(row=start, column=4, value="미확인 항목이 많은 메모 (상위 8건)").font = SECTION_FONT
    _header(ws, start + 1, ["메모", "긴급도", "미확인"], col=4)
    ranking = unconfirmed_ranking(records, 8)
    for i, (name, priority, count) in enumerate(ranking, start=start + 2):
        ws.cell(row=i, column=4, value=name)
        ws.cell(row=i, column=5, value=priority).font = Font(bold=True, color=PRIORITY_INK[priority])
        ws.cell(row=i, column=6, value=count)
    if ranking:
        bar = BarChart()
        bar.type, bar.title, bar.legend = "bar", "미확인 항목이 많은 메모", None
        first, last = start + 2, start + 1 + len(ranking)
        bar.add_data(Reference(ws, min_col=6, min_row=start + 1, max_row=last), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=4, min_row=first, max_row=last))
        bar.y_axis.majorGridlines = None
        bar.x_axis.scaling.orientation = "maxMin"  # 많은 순이 위로
        for i, (_, priority, _) in enumerate(ranking):
            point = DataPoint(idx=i)
            point.graphicalProperties.solidFill = STATUS_BAR[priority]
            bar.series[0].dPt.append(point)
        bar.height, bar.width = 8, 12
        ws.add_chart(bar, f"H{start}")
    _widths(ws, {"A": 16, "B": 44, "C": 4, "D": 30, "E": 10, "F": 10})


def _detail_rows(record: Record) -> list[list]:
    r = record.report
    base = [record.source_name, r.priority]
    rows = [base + ["요약", r.summary, ""], base + ["긴급도 판단 이유", r.priority_reason, ""],
            base + ["문제 유형", r.category, ""]]
    groups = [
        ("발생한 일", r.events, lambda i: i.text),
        ("조치한 내용", r.actions, lambda i: i.text),
        ("현재 상태", r.current_status, lambda i: i.text),
        ("다음 근무자 확인 사항", r.next_checks, lambda i: i.text),
        ("시간순 경과", r.timeline, lambda t: f"{t.time}  {t.text}"),
        ("주요 수치·대상", r.key_facts, lambda f: f"{f.label}: {f.value}"),
        ("연락·요청 현황", r.contacts, lambda c: f"{c.party} — {c.request} — {c.status}"),
    ]
    for title, items, fmt in groups:
        rows += [base + [title, fmt(item), item.source_quote] for item in items]
    rows += [base + ["미확인 사항", f"{u.subject}: 미확인" + (f" ({u.note})" if u.note else ""), ""]
             for u in r.unconfirmed]
    return rows


def _insight_rows(snapshot: InsightSnapshot | None) -> list[list]:
    if snapshot is None:
        return [["", "분석한 인사이트가 없습니다", "대시보드에서 ‘다시 분석’을 실행하세요", ""]]
    ins = snapshot.insights
    rows = [["종합", ins.headline, f"보고서 {snapshot.source_count}건 기준 · {snapshot.created_at:%Y-%m-%d %H:%M}", ""]]
    for label, points in (("공통 패턴", ins.patterns), ("주의할 경향", ins.risks), ("기록 보완점", ins.gaps)):
        rows += [[label, p.title, p.detail, ", ".join(p.related)] for p in points]
    return rows


def export_workbook(records: list[Record], snapshot: InsightSnapshot | None, path: Path,
                    now: datetime | None = None) -> Path:
    now = now or datetime.now()
    records = sort_for_dashboard(records)
    wb = Workbook()
    dash = wb.active
    dash.title = SHEETS[0]
    _dashboard(dash, records, snapshot, now)

    _table(wb.create_sheet(SHEETS[1]),
           ["긴급도", "메모", "문제 유형", "요약", "판단 이유", "미확인", "경고", "처리 시각", "보고서 파일"],
           [[r.report.priority, r.source_name, r.report.category, r.report.summary, r.report.priority_reason,
             len(r.report.unconfirmed), len(r.warnings), r.created_at.strftime("%Y-%m-%d %H:%M"), r.txt_path.name]
            for r in records],
           {"A": 8, "B": 28, "C": 14, "D": 60, "E": 60, "F": 8, "G": 8, "H": 17, "I": 48}, priority_col=1)

    _table(wb.create_sheet(SHEETS[2]), ["메모", "긴급도", "항목", "내용", "근거 문장"],
           [row for r in records for row in _detail_rows(r)],
           {"A": 28, "B": 8, "C": 20, "D": 70, "E": 60}, priority_col=2)

    _table(wb.create_sheet(SHEETS[3]), ["구분", "제목", "설명", "관련 메모"], _insight_rows(snapshot),
           {"A": 12, "B": 34, "C": 80, "D": 50})

    _table(wb.create_sheet(SHEETS[4]), ["메모", "긴급도", "경고"],
           [[r.source_name, r.report.priority, w] for r in records for w in r.warnings],
           {"A": 28, "B": 8, "C": 100}, priority_col=2)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
